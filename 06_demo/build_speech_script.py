from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(r"C:\Users\admin\Documents\Codex\2026-09-17\pe6201-group-project-continuation\deliverables\AI_Scroll_5_Minute_Speech_Script.docx")

NAVY = "0B2D49"
TEAL = "139F99"
PALE = "EAF7F5"
LIGHT = "F3F7FA"
MUTED = "60748A"
CORAL = "FF7474"
WHITE = "FFFFFF"

SLIDES = [
    {
        "title": "Slide 1 — Introduction",
        "time": "0:00–0:36 · 36 seconds",
        "cue": "Show the cover slide. Pause after the project name.",
        "en": "Hi, I’m Jason. My project is called AI Scroll. It is an academic email agent for students who receive deadlines, extensions and course updates across many messages. The system reads selected course emails, keeps the original evidence, merges related updates and prepares a practical study plan. It can then create calendar proposals, but only after the student reviews and approves them. In this presentation, I will explain the problem, the Agent design, the tools I built, and the current working result.",
        "zh": "大家好，我是 Jason。我的项目叫 AI Scroll，是一个面向学生的学术邮件 Agent。它读取指定课程的邮件，保留原始证据，合并相关更新，并生成学习计划。只有在学生检查并批准后，它才会提出日历写入建议。接下来我会介绍问题、Agent 设计、工具和当前成果。",
    },
    {
        "title": "Slide 2 — Customer pain points",
        "time": "0:36–1:19 · 43 seconds",
        "cue": "Emphasise “one reliable decision” and “preserve evidence”.",
        "en": "The main problem is not a lack of email. It is the effort required to turn several messages into one reliable decision. A deadline can appear in different formats, and a later message may extend or cancel it. Students also receive reminders, newsletters and platform notifications that look important but do not create a real deadline. If the user manually copies every date, errors are easy. If an Agent writes to a calendar automatically, the risk is even higher. So the product must understand the message chain and preserve evidence before it takes any external action.",
        "zh": "真正的问题不是邮件太少，而是学生要把多封邮件整理成一个可靠的最终决定。截止日期可能使用不同格式，之后的邮件还可能延期或取消。系统必须理解邮件链，并在执行外部操作前保存证据。",
    },
    {
        "title": "Slide 3 — Product flow",
        "time": "1:19–2:01 · 42 seconds",
        "cue": "Point from left to right as you describe the six steps.",
        "en": "AI Scroll follows six steps. First, a mailbox connector reads new messages. Second, a course scope filter removes unrelated mail before any model call. Third, the extraction layer identifies the task, deadline and exact evidence sentence. Fourth, the merge engine combines related emails and keeps the latest valid state. Fifth, the planner works backwards from the verified deadline to create manageable sessions. Finally, the approval gate shows the proposal to the student. The Agent prepares the plan, but the student owns the final calendar action. This separation is important for both safety and trust.",
        "zh": "AI Scroll 有六步：读取、筛选、提取、合并、倒推计划和人工批准。Agent 负责准备决定，学生保留最终日历操作权，因此兼顾安全和信任。",
    },
    {
        "title": "Slide 4 — Agent and tools",
        "time": "2:01–2:45 · 44 seconds",
        "cue": "Follow Observe → Reason → Act, then name the five tools.",
        "en": "The system uses a simple ReAct style loop: observe, reason and act. The Agent observes new messages and the existing timeline. It reasons about which tools are needed and whether a message is new, related or irrelevant. It then acts by preparing a timeline item or a calendar proposal, before stopping for review. The main tools read the mailbox, filter course messages, extract structured JSON, match and merge related threads, and enforce the calendar approval gate. I use GPT 5 Mini through OpenRouter for language understanding. Deterministic rules still validate dates, calculate urgency, remove duplicates and block unsafe writes.",
        "zh": "系统采用 Observe、Reason、Act 的 ReAct 循环。五个工具分别负责读取邮箱、筛选课程邮件、提取结构化 JSON、匹配合并邮件线程，以及强制执行日历审批。语言理解使用 GPT 5 Mini，日期、紧急度、去重和审批仍由确定性规则控制。",
    },
    {
        "title": "Slide 5 — Hybrid architecture",
        "time": "2:45–3:26 · 41 seconds",
        "cue": "Compare the two columns: AI flexibility and rule control.",
        "en": "I chose a hybrid architecture because the two parts of the problem need different strengths. The language model handles varied academic language, extracts evidence and interprets whether a message is a new task or an update. The application code handles the parts that must remain predictable: date validation, Singapore time, urgency thresholds, thread history, deduplication and approval. I built the workflow, merge logic, evaluation and safety controls. I rent the language model and the external mailbox and calendar APIs. This reduces development time while keeping the important project logic visible and testable.",
        "zh": "混合架构让语言模型处理灵活的自然语言，让程序规则处理必须稳定的日期验证、时区、紧急度、历史记录、去重和审批。我自行构建工作流、合并逻辑、评估和安全控制，外部租用模型、邮箱和日历 API。",
    },
    {
        "title": "Slide 6 — Final product experience",
        "time": "3:26–4:12 · 46 seconds",
        "cue": "Show the real interface. Trace inbox → evidence → action plan → review.",
        "en": "This is the current product experience. The left side behaves like an academic inbox. The middle panel keeps the selected email visible, including the exact sentence that created or changed the deadline. The right panel turns the message into a five step action plan. It shows the final deadline, the planned stages and the evidence used by the Agent. The scan button refreshes the interface using the actual mailbox pipeline. Settings stay in a collapsed sidebar, so the main screen remains simple. The user then opens Calendar Review, checks the proposal and confirms it. Without that confirmation, the system performs no calendar write.",
        "zh": "当前界面左侧是学术邮箱，中间保留原邮件和截止日期证据，右侧生成五步行动计划。扫描按钮调用真实邮箱流程；用户进入 Calendar Review 检查并确认，未确认时系统不会写入日历。",
    },
    {
        "title": "Slide 7 — Evidence and close",
        "time": "4:12–5:01 · 49 seconds",
        "cue": "State the draft-label limitation clearly, then close on student control.",
        "en": "The current evidence shows that the pipeline improves on a simple keyword and date baseline, especially for calendar actions and cross email merging. These comparison numbers still use draft labels, so I will not present them as the final assessment result until the 50 labels are reviewed and frozen. The completed live model check passed all six web derived edge cases, the code passes 47 automated tests, and the six case live run cost about 1.1 cents. The next steps are to freeze the labels, run the final evaluation and record the live demo. AI Scroll shows how an Agent can reduce email overload while keeping the student in control.",
        "zh": "当前结果优于简单关键词和日期基线，尤其是在日历动作和跨邮件合并方面。但比较数据仍使用草稿标签，必须复核并冻结 50 个标签后才能作为最终评估。当前六个真实网页案例全部通过，代码通过 47 项自动化测试，六案例实时运行成本约 1.1 美分。",
    },
]


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=None, font="Arial"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_label(doc, text, color=TEAL):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper())
    set_run_font(r, 9, True, color)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(NAVY)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

for style_name, size, color in (("Title", 32, NAVY), ("Heading 1", 21, NAVY), ("Heading 2", 14, TEAL)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)

if "Cue" not in styles:
    cue_style = styles.add_style("Cue", WD_STYLE_TYPE.PARAGRAPH)
else:
    cue_style = styles["Cue"]
cue_style.font.name = "Arial"
cue_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
cue_style.font.size = Pt(9.5)
cue_style.font.italic = True
cue_style.font.color.rgb = RGBColor.from_string(MUTED)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(40)
r = p.add_run("PE6201 COURSE PROJECT")
set_run_font(r, 11, True, TEAL)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(16)
p.paragraph_format.space_after = Pt(6)
r = p.add_run("AI Scroll")
set_run_font(r, 34, True, NAVY)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(22)
r = p.add_run("Five-Minute Presentation Script")
set_run_font(r, 22, True, TEAL)

box = doc.add_table(rows=1, cols=1)
box.autofit = False
box.columns[0].width = Inches(6.6)
cell = box.cell(0, 0)
set_cell_fill(cell, PALE)
set_cell_margins(cell, 260, 260, 260, 260)
for idx, line in enumerate([
    "Speaker  Jason Wang Chen Yu",
    "Target  5:00–5:20",
    "Hard limit  Keep the final recording under 8:00",
    "Format  Slides + product UI demonstration",
]):
    p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    label, value = line.split("  ", 1)
    r = p.add_run(label + "  ")
    set_run_font(r, 10.5, True, NAVY)
    r = p.add_run(value)
    set_run_font(r, 10.5, False, MUTED)

doc.add_paragraph("")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Evidence first  ·  Human approval before calendar writes")
set_run_font(r, 11, True, TEAL)

doc.add_page_break()

# Run of show
doc.add_heading("Recording plan", level=1)
p = doc.add_paragraph("A natural delivery of this 676-word script is approximately five minutes. Use the slide changes as breathing points and keep the live interface explanation focused on the visible workflow.")
p.paragraph_format.space_after = Pt(12)

table = doc.add_table(rows=1, cols=3)
table.autofit = False
widths = [Inches(0.65), Inches(3.75), Inches(1.65)]
for i, (label, width) in enumerate(zip(["Slide", "Purpose", "Time"], widths)):
    table.columns[i].width = width
    c = table.rows[0].cells[i]
    set_cell_fill(c, NAVY)
    set_cell_margins(c)
    r = c.paragraphs[0].add_run(label)
    set_run_font(r, 9.5, True, WHITE)

purposes = [
    "Product introduction", "Customer pain points", "Six-step product flow",
    "ReAct loop and tools", "Hybrid architecture", "Current product experience",
    "Evidence, limitations and close",
]
for i, (slide, purpose) in enumerate(zip(SLIDES, purposes), 1):
    cells = table.add_row().cells
    fill = LIGHT if i % 2 else WHITE
    for c in cells:
        set_cell_fill(c, fill)
        set_cell_margins(c, 100, 120, 100, 120)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for c, value in zip(cells, [str(i), purpose, slide["time"].split(" · ")[0]]):
        r = c.paragraphs[0].add_run(value)
        set_run_font(r, 9.5, i == 1, NAVY)

doc.add_paragraph("")
add_label(doc, "Recording reminders")
for item in [
    "Speak at 130–140 words per minute. Do not rush the metrics on Slide 7.",
    "Keep the mouse still unless you are pointing to a visible product element.",
    "If time runs long, shorten the Chinese preparation notes; do not read them in the English recording.",
]:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(item)
    set_run_font(r, 10.5, False, NAVY)

doc.add_page_break()

# Script sections
for index, slide in enumerate(SLIDES, 1):
    doc.add_heading(slide["title"], level=1)
    p = doc.add_paragraph()
    r = p.add_run(slide["time"])
    set_run_font(r, 10, True, TEAL)

    p = doc.add_paragraph(style="Cue")
    r = p.add_run("SLIDE CUE  " + slide["cue"])
    set_run_font(r, 9.5, False, MUTED)

    add_label(doc, "English script")
    p = doc.add_paragraph(slide["en"])
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(8)

    add_label(doc, "中文理解", CORAL)
    p = doc.add_paragraph(slide["zh"])
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(12)

    if index in (2, 4, 6):
        doc.add_page_break()
    elif index != len(SLIDES):
        sep = doc.add_paragraph()
        sep.paragraph_format.space_before = Pt(3)
        sep.paragraph_format.space_after = Pt(3)
        run = sep.add_run("—" * 26)
        set_run_font(run, 8, False, "C9D8E3")

doc.add_page_break()
doc.add_heading("Pronunciation and final checklist", level=1)

pron = [
    ("academic", "ak-uh-DEM-ik", "学术的"),
    ("evidence", "EV-i-duhns", "证据"),
    ("deterministic", "di-TUR-mi-NIS-tik", "确定性的"),
    ("deduplication", "dee-doo-pli-KAY-shun", "去重"),
    ("provisional", "pruh-VIZH-uh-nuhl", "暂定的"),
    ("urgency", "UR-juhn-see", "紧急程度"),
    ("OpenRouter", "OH-puhn ROW-ter", "模型路由服务"),
]
table = doc.add_table(rows=1, cols=3)
for i, label in enumerate(["Term", "Suggested reading", "Meaning"]):
    c = table.rows[0].cells[i]
    set_cell_fill(c, NAVY)
    r = c.paragraphs[0].add_run(label)
    set_run_font(r, 9.5, True, WHITE)
for idx, row in enumerate(pron, 1):
    cells = table.add_row().cells
    for c, value in zip(cells, row):
        set_cell_fill(c, LIGHT if idx % 2 else WHITE)
        set_cell_margins(c, 90, 120, 90, 120)
        r = c.paragraphs[0].add_run(value)
        set_run_font(r, 9.5, False, NAVY)

doc.add_paragraph("")
add_label(doc, "Before recording")
checks = [
    "Open the final PPT and the local AI Scroll interface before starting Zoom.",
    "Run one mailbox scan and keep a successful example ready for Slide 6.",
    "Use 16:9 screen sharing and record at readable browser zoom.",
    "Keep the final video close to five minutes and below the teacher’s eight-minute viewing limit.",
    "Export once, replay the full video, and confirm that voice, slides and interface text are clear.",
]
for item in checks:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(item)
    set_run_font(r, 10.5, False, NAVY)

# Header/footer
for sec in doc.sections:
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("AI Scroll · PE6201")
    set_run_font(r, 8, True, TEAL)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Jason Wang Chen Yu  |  Five-Minute Presentation Script")
    set_run_font(r, 8, False, MUTED)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
