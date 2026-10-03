from __future__ import annotations

from pathlib import Path
import subprocess

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "05_report" / "AI_Scroll_Trade_Off_Analysis_Final.docx"
ASSET_DIR = ROOT / "05_report" / "assets"
UI_IMAGE = ROOT / "06_demo" / "ui_final_wide.png"
CHART_IMAGE = ASSET_DIR / "evaluation_comparison.png"

NAVY, TEAL, MINT = "0B2D49", "119E96", "E6F6F3"
PALE, MUTED, WHITE = "F3F7FA", "60748A", "FFFFFF"


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), color)


def margins(cell, top=110, start=140, bottom=110, end=140) -> None:
    props = cell._tc.get_or_add_tcPr()
    tc_mar = props.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        props.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def font(run, size=10, bold=False, color=NAVY) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    font(p.add_run(text), 10)


def heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run(text), 15, True)


def callout(doc: Document, label: str, text: str, fill=MINT) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    margins(cell, 180, 190, 180, 190)
    p = cell.paragraphs[0]
    font(p.add_run(label.upper() + "\n"), 9, True, TEAL)
    font(p.add_run(text), 11, True)


def compact_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (label, width) in enumerate(zip(headers, widths)):
        table.columns[i].width = Inches(width)
        cell = table.rows[0].cells[i]
        shade(cell, NAVY)
        margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(label), 8.5, True, WHITE)
    for row_i, values in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cell, WHITE if row_i % 2 == 0 else PALE)
            margins(cell, 90, 115, 90, 115)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            font(p.add_run(value), 8.5, i == 0)


def make_chart() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    subprocess.run([r"E:\Python\python.exe", str(ROOT / "05_report" / "build_evaluation_chart.py")], check=True)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(paragraph.add_run("AI Scroll  ·  PE6201  |  "), 8, False, MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def build() -> None:
    make_chart()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.62)
    sec.left_margin = sec.right_margin = Inches(0.72)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    font(p.add_run("PE6201 · END-OF-COURSE PROJECT"), 9, True, TEAL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    font(p.add_run("AI Scroll"), 28, True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    font(p.add_run("Business and Technical Trade-off Analysis"), 18, True, TEAL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    font(p.add_run("Wang Chen Yu Jason  ·  Submission version  ·  4 October 2026"), 9, True, MUTED)
    callout(doc, "Executive decision", "Use an evidence-first hybrid Agent: a language model interprets academic email, deterministic controls validate and merge results, and the student approves every calendar write.")
    heading(doc, "Problem and product decision")
    add_body(doc, "Students receive deadlines, extensions, cancellations and submission rules across many messages. The hard case is a chain in which a later email changes the original instruction. Manual search and copy-paste are slow and can preserve an obsolete deadline. AI Scroll converts selected course emails into one evidence-backed timeline, a backward study plan and a calendar proposal. The system automates interpretation and preparation while keeping the final external action under explicit user control.")
    doc.add_picture(str(UI_IMAGE), width=Inches(7.0))
    cap = doc.add_paragraph("Figure 1. Working interface: inbox evidence, selected message and approval-gated action plan.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(cap.runs[0], 8, False, MUTED)

    doc.add_page_break()
    heading(doc, "System scope and hybrid architecture")
    add_body(doc, "The Agent reads from authorized Gmail or Outlook accounts, filters by course before a model call, extracts a structured task and exact evidence sentence, matches related messages, applies later corrections, and plans backwards from the verified deadline. GPT-5 Mini through OpenRouter handles varied language. Application code controls ISO date validation, Singapore time, urgency thresholds, deduplication, merge history and approval. Google Calendar writes are available only after review.")
    compact_table(doc, ["Decision", "Selected approach", "Benefit", "Cost or constraint"], [
        ["Language vs rules", "Hybrid Agent", "Flexible interpretation with inspectable controls", "More integration and validation code"],
        ["Build vs buy", "Build workflow; rent model and APIs", "Fast development; replaceable model", "Provider cost, latency and privacy exposure"],
        ["Mailbox access", "Narrow OAuth connectors", "No manual email copy-paste", "Authorization and tenant-policy dependency"],
        ["Calendar action", "Preview then approve", "Blocks silent or unsafe writes", "One extra user step"],
        ["Polling", "Metadata first; model only for changes", "Lower cost at high freshness", "Requires message-state tracking"],
    ], [1.2, 1.55, 2.25, 2.05])
    heading(doc, "Business and operating trade-offs")
    add_body(doc, "I build the workflow, schemas, thread matcher, merge engine, urgency policy, evaluation harness and approval gate. I rent the language model and use existing mailbox and calendar infrastructure. This avoids hosting model weights and shortens development time, while configuration keeps the model endpoint replaceable. The trade-off is dependency on provider availability, dated prices and OAuth policy.")
    add_body(doc, "Freshness does not require paying for a model call every time the inbox is checked. A naive 15-minute model job creates 2,880 monthly calls even when nothing changes; hourly polling creates 720. AI Scroll checks identifiers and course scope first, then calls the model only for new or changed messages. At 20 relevant emails per day, both schedules require about 600 model calls per 30-day month.")
    callout(doc, "Measured live cost", "Six web-derived cases cost US$0.01074 in total, or about US$0.00179 per processed case.", fill=PALE)

    doc.add_page_break()
    heading(doc, "Data and evaluation design")
    add_body(doc, "The development set contains 50 synthetic academic emails across 38 threads: 30 independent emails and eight multi-email threads containing 20 messages. It includes 11 updates, three clarification cases and one cancellation. Each row contains draft labels for course, task type, title, deadline, urgency, evidence, relation and calendar action. Metrics are reported separately because task classification, deadline extraction, urgency and merging are different problems.")
    doc.add_picture(str(CHART_IMAGE), width=Inches(7.0))
    cap = doc.add_paragraph("Figure 2. Development comparison using draft labels; percentages use the eligible cases for each capability.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(cap.runs[0], 8, False, MUTED)
    compact_table(doc, ["Evidence item", "Result", "Interpretation"], [
        ["Automated tests", "47 passed", "Core extraction, merge, planning and safety paths"],
        ["Live web-derived cases", "6 / 6 passed", "Current model handled the selected edge cases"],
        ["Cross-email final state", "8 / 8", "All development chains reached the expected latest state"],
        ["Unauthorized calendar writes", "0", "Approval gate blocked unconfirmed external action"],
    ], [2.1, 1.3, 3.55])
    add_body(doc, "The pipeline improves strongly on the keyword-and-date baseline for task type, calendar action and cross-email merging. These figures support the design direction, but they remain development evidence because the labels are not yet frozen and the data is synthetic. The evaluation preserves the denominator for each capability and avoids presenting one blended accuracy score.")

    doc.add_page_break()
    heading(doc, "Risks and implemented controls")
    compact_table(doc, ["Risk", "Implemented control", "Remaining limitation"], [
        ["Wrong or obsolete deadline", "Exact evidence, thread history and later-update merge", "Ambiguous language can still require clarification"],
        ["Unsafe calendar action", "Preview, expiry check, deduplication and human approval", "User may approve an incorrect proposal"],
        ["Privacy exposure", "Course filtering, narrow OAuth and minimal structured storage", "External providers still process selected content"],
        ["Provider failure", "Replaceable model configuration and deterministic fallback paths", "Quality may change across providers"],
        ["Overstated evaluation", "Capability counts, fixed denominators and draft-label disclosure", "Synthetic data cannot prove generalization"],
    ], [1.65, 3.1, 2.2])
    heading(doc, "Limitations and next evidence")
    add_body(doc, "The dataset is synthetic, its holdout was generated by the same process, and only eight threads test cross-email reasoning. Six live cases are useful for integration testing but too small for a performance claim. A final evaluation should freeze human-reviewed labels, run one fixed model and prompt over all cases, preserve raw outputs and token logs, calculate dated cost per email and thread, and report residual errors without tuning on the test split.")
    heading(doc, "Conclusion")
    add_body(doc, "AI Scroll demonstrates a practical boundary for an academic email Agent. The model handles language, while deterministic code protects dates, history, duplication and approval. The product reduces manual consolidation without removing student control. The current prototype is strong enough to demonstrate the end-to-end workflow; broader real-email evaluation and privacy review are required before deployment.")
    callout(doc, "Recommendation", "Submit the current prototype as an evidence-backed course project and position real-world deployment as the next validated stage.")

    for section in doc.sections:
        add_page_number(section.footer.paragraphs[0])
    doc.core_properties.title = "AI Scroll Business and Technical Trade-off Analysis"
    doc.core_properties.author = "Wang Chen Yu Jason"
    doc.core_properties.subject = "PE6201 End-of-Course Project"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
